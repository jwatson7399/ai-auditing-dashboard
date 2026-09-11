#!/usr/bin/env python3
"""Daily rebuild for the AI Auditing dashboard.

Runs on GitHub Actions (see .github/workflows/rebuild.yml). Every number on the
page comes from this script: structured API responses, the merged inbox on main,
and the run logs. No language model is involved in any step here.

    python tools/rebuild.py                 # full run, needs AA_API_KEY
    python tools/rebuild.py --no-fetch      # rebuild from the last data file only
    python tools/rebuild.py --probe         # print what the sources return, write nothing
    python tools/rebuild.py --no-log        # do not append to log/runs-rebuild.csv

Outputs: data/YYYY-MM-DD.json, data/latest.json, site/index.html,
log/run-log.xlsx, one line in log/runs-rebuild.csv.

Status values (same meanings as the agents):
  ok       every source fetched fresh, page built
  partial  page built but one or more sources carried forward from the last run
  blocked  a required credential or input was missing, page not rebuilt
  error    an unexpected failure, page not rebuilt
"""
import argparse
import csv
import datetime as dt
import glob
import json
import os
import re
import sys
import traceback
import zoneinfo
from public_sources import (AA_URL, AA_COLUMNS, ARENA_URLS, fetch_public_table,
                            parse_aa_table, parse_arena_table, effort_key)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ET = zoneinfo.ZoneInfo("America/New_York")
AGENTS = ["benchmark-scout", "news-scout", "editor", "verifier", "commentator", "rebuild"]
AGENT_LABEL = {"benchmark-scout": "Benchmark Scout", "news-scout": "News Scout", "editor": "Editor",
               "verifier": "Verifier", "commentator": "Commentator", "rebuild": "Daily rebuild"}

# --- Access: the only models the page may recommend. Human-maintained. ---------
ACCESS = {
    "Claude Fable 5.1": "claude", "Claude Fable 5": "claude", "Claude Opus 5": "claude",
    "GPT-6 Astra": "chatgpt", "GPT-5.6 Sol": "chatgpt", "GPT-5.6 Terra": "chatgpt", "GPT-5.6 Luna": "chatgpt",
    "Grok 4.6": "grok",
}
MEDIA_ACCESS = {"GPT Image 2": "chatgpt", "GPT Image 1.5": "chatgpt", "Grok Imagine": "grok", "Grok Imagine Video": "grok"}

# --- Benchmarks and the weight table (strategy.md, phase one weights) -----------
# api: candidate field names in the Artificial Analysis evaluations object, in order of
# preference. The first one present is used; the field actually used is stored with the
# scores. If none is present the benchmark is carried forward and the run is partial.
BENCH = {
    # One candidate only, deliberately. first_key is resolved per model, so a fallback would
    # let some models contribute terminalbench_hard into a column headed Terminal-Bench v2.1.
    # terminalbench_hard is the older test and is null on every current model; v2.1 is the one
    # this page has always claimed. If AA renames the field the benchmark carries forward stale,
    # which is the intended failure.
    "terminal": {"name": "Terminal-Bench v2.1", "kind": "task", "api": ["terminalbench_v2_1"]},
    "scicode":  {"name": "SciCode", "kind": "task", "api": ["scicode"]},
    "gdpval":   {"name": "GDPval-AA v2", "kind": "task", "api": ["gdpval_aa", "gdpval"]},
    # tau_banking is AA's tau-cubed Bench Banking; the API just omits the superscript. Verified
    # against the published figures: Fable 5.1 max 0.472 and Astra max 0.414 match 47 and 41.
    # tau2 is the older tau-squared Bench and must not be substituted, so no fallback here.
    "tau":      {"name": "τ³-Banking", "kind": "task", "api": ["tau_banking"]},
    "lcr":      {"name": "AA-LCR", "kind": "task", "api": ["lcr", "aa_lcr"]},
    "omni":     {"name": "AA-Omniscience Accuracy", "kind": "task", "api": ["aa_omniscience_accuracy", "omniscience_accuracy"]},
    "nohalluc": {"name": "Non-hallucination rate", "kind": "task", "api": ["aa_omniscience_non_hallucination_rate", "non_hallucination_rate"],
                 "derive_from_hallucination": ["aa_omniscience_hallucination_rate", "hallucination_rate"]},
    "gdppdf":   {"name": "GDP.pdf (all criteria met)", "kind": "task", "api": ["gdp_pdf", "gdppdf"]},
    "webdev":   {"name": "Arena WebDev", "kind": "votes", "elo": True, "arena": "webdev"},
    "text":     {"name": "Arena Text", "kind": "votes", "elo": True, "arena": "text"},
    "image":    {"name": "AA Image Arena", "kind": "votes", "elo": True, "media": "text-to-image"},
    # Named for what the endpoint returns, which is the whole text-to-video board. It is not the
    # narrower "with audio" chart the seed numbers were read from: the response comes back with
    # include_categories false and every row's categories list empty, so no filter is applied,
    # and it carries 83 models where that chart had a handful. Calling it the audio board made
    # a change of board look like every model gaining about 95 Elo overnight.
    "video":    {"name": "AA Video Arena", "kind": "votes", "elo": True, "media": "text-to-video"},
}
INDEX_API = ["artificial_analysis_intelligence_index"]
# Token price comes from the API; per-task cost is joined separately from the public table.
PRICE_API = ["price_1m_blended_3_to_1"]
PRICE_LABEL = "Price per 1M tokens, 3:1 blend"
# An Elo board outside this range means the cell we read is not the score. Raising carries the
# board forward instead of publishing a wrong number.
ELO_RANGE = (800, 2000)

TASKS = [
    {"id": "coding", "name": "Coding projects", "w": {"terminal": 0.6, "scicode": 0.4},
     "tests": "Terminal-Bench hands the model a command line and a real job (install this, fix that build, process these files) and checks whether it finishes. SciCode has it write working code for research problems.",
     "daily": "\"Can it do technical chores on a computer by itself, and is the code it writes correct?\""},
    {"id": "agents", "name": "Supervising coding agents", "w": {"gdpval": 0.3, "tau": 0.25, "terminal": 0.25, "lcr": 0.2}, "extra": "gdppdf",
     "tests": "Blends the long, multi-step tests: real job deliverables judged against other models (GDPval), following rules while using tools across a long conversation (τ³-Banking), terminal work, and reasoning over long documents.",
     "daily": "\"When I hand it a big project and walk away, will it keep track, follow instructions, and not go off the rails?\""},
    {"id": "frontend", "name": "Frontend and design", "w": {"webdev": 0.7, "terminal": 0.3},
     "tests": "Arena WebDev shows two models the same request, builds both sites, and asks people which they prefer. Terminal-Bench adds whether it can actually ship the code.",
     "daily": "\"Which one makes the page people like more, and can it build it?\" Mostly a taste vote."},
    {"id": "image", "name": "Image generation", "w": {"image": 1}, "media": True,
     "tests": "Blind human votes between two images made from the same prompt.",
     "daily": "\"Which image tool do people prefer?\" Pure preference, no task test exists."},
    {"id": "video", "name": "Video generation", "w": {"video": 1}, "media": True,
     "tests": "Blind human votes between two clips made from the same prompt, across the whole video board rather than the audio-only chart.",
     "daily": "\"Which video tool do people prefer?\" Pure preference. Sora is not on this board at all."},
    {"id": "writing", "name": "Writing to humans", "w": {"text": 0.5, "nohalluc": 0.25, "omni": 0.25},
     "tests": "Arena Text is a blind vote on which reply people prefer in ordinary chat. Non-hallucination is how often the model admits it does not know instead of bluffing. Accuracy is plain factual recall.",
     "daily": "\"Which one writes the way people like, and will it quietly make things up in my email?\""},
    {"id": "audit", "name": "Auditing code", "w": {"terminal": 0.4, "scicode": 0.3, "nohalluc": 0.3},
     "tests": "Coding tests plus the bluffing rate, because a reviewer that invents bugs is worse than none.",
     "daily": "\"Will it find the real problems in this code without inventing fake ones?\""},
    {"id": "technical", "name": "Technical jobs", "w": {"terminal": 0.4, "gdpval": 0.3, "lcr": 0.3}, "extra": "gdppdf",
     "tests": "Terminal work, real job deliverables, and long-document reasoning.",
     "daily": "\"Data cleanup, scripts, configuring things, reading a long spec: can it handle the whole job?\""},
]
NOISE = 3  # points; gaps under this are "within noise"


class Blocked(Exception):
    pass


def now_et():
    return dt.datetime.now(ET)


def log(msg):
    print(msg, flush=True)


# ------------------------------------------------------------------ data files
def data_files():
    return sorted(glob.glob(os.path.join(ROOT, "data", "????-??-??.json")))


def load_previous(today):
    prev = [f for f in data_files() if os.path.basename(f)[:10] < today]
    if not prev:
        return None
    with open(prev[-1]) as f:
        return json.load(f)


# ------------------------------------------------------------ Artificial Analysis
def http_get(url, headers=None, timeout=60):
    import requests
    r = requests.get(url, headers=headers or {}, timeout=timeout)
    return r.status_code, r


def split_effort(name):
    m = re.match(r"^(.*?)\s*\(([^)]*)\)\s*$", name)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return name.strip(), "default"


# Effort settings appear in three shapes across the three endpoints: a parenthesis on the llms
# and image boards, a trailing word on the arena slugs, and nothing at all. Stripped only from
# the end of a name, so a model whose own name contains one of these words keeps it.
EFFORT_TOKENS = ("max", "xhigh", "high", "medium", "low", "thinking", "reasoning", "non")


def canon(name):
    """A comparison key for one model, ignoring case, punctuation and effort setting."""
    base, _ = split_effort(name)
    toks = re.sub(r"[^a-z0-9]+", " ", base.lower()).split()
    while toks and toks[-1] in EFFORT_TOKENS:
        toks.pop()
    return " ".join(toks)


def fold_to_access(d, access):
    """Fold the names an endpoint returns onto the display names in the access table.

    The arena and media endpoints return slugs (grok-imagine-video) and effort suffixes
    (GPT Image 2 (high)) where the llms endpoint returns display names, so without this the
    access lookup misses and the task has no pick. Same rule as the llms path: one row per
    model at its best effort setting. Every fold is returned so the day's JSON can show it,
    and a name that matches nothing is kept verbatim rather than dropped, so an unrecognised
    model is visible on the page instead of silently disappearing.
    """
    lookup = {canon(k): k for k in access}
    folded, aliases, unmatched = {}, {}, []
    for raw, val in d.items():
        base = lookup.get(canon(raw))
        if base is None:
            unmatched.append(raw)
            if raw not in folded or val > folded[raw]:
                folded[raw] = val
            continue
        aliases[raw] = base
        if base not in folded or val > folded[base]:
            folded[base] = val
    return folded, aliases, sorted(unmatched)


def first_key(d, candidates):
    for k in candidates:
        if k in d and d[k] is not None:
            return k
    return None


def as_percent(v):
    return round(v * 100, 1) if 0 <= v <= 1 else round(v, 1)


def fetch_aa_llms(key, probe=False):
    """Returns (bench_partial, eff, index_meta, raw). Raises Blocked without a key."""
    if not key:
        raise Blocked("AA_API_KEY not set")
    status, r = http_get("https://artificialanalysis.ai/api/v2/data/llms/models", {"x-api-key": key})
    if status != 200:
        raise RuntimeError(f"AA llms endpoint returned {status}: {r.text[:200]}")
    raw = r.json()
    models = raw.get("data", raw if isinstance(raw, list) else [])
    if probe:
        log(f"AA llms: {len(models)} models")
        if models:
            log("top-level keys: " + ", ".join(sorted(models[0].keys())))
            log("evaluation keys: " + ", ".join(sorted((models[0].get("evaluations") or {}).keys())))
            log("names: " + " | ".join(m.get("name", "?") for m in models[:60]))
        return None
    bench, eff = {}, []
    used_fields = {}
    for m in models:
        name = m.get("name") or m.get("slug") or ""
        base, effort = split_effort(name)
        ev = m.get("evaluations") or {}
        idx_k = first_key(ev, INDEX_API)
        idx = ev.get(idx_k) if idx_k else None
        pricing = m.get("pricing") or {}
        price_k = first_key(pricing, PRICE_API)
        price = pricing.get(price_k) if price_k else None
        if idx is not None:
            eff.append({"model": base, "effort": effort, "score": round(idx, 1), "price_1m": price,
                        "speed": m.get("median_output_tokens_per_second"),
                        "wait_s": m.get("median_time_to_first_answer_token") or m.get("median_time_to_first_token_seconds"),
                        "access": ACCESS.get(base), "estimate": False, "api_name": name})
        for k, b in BENCH.items():
            if "api" not in b:
                continue
            fk = first_key(ev, b["api"])
            val = None
            if fk:
                val = as_percent(ev[fk])
            elif b.get("derive_from_hallucination"):
                hk = first_key(ev, b["derive_from_hallucination"])
                if hk:
                    val = round(100 - as_percent(ev[hk]), 1)
                    fk = hk + " (100 minus)"
            if val is None:
                continue
            used_fields[k] = fk
            slot = bench.setdefault(k, {"d": {}, "settings": {}})
            # Benchmark charts use each model's best effort setting; the setting is stored.
            if base not in slot["d"] or val > slot["d"][base]:
                slot["d"][base] = val
                slot["settings"][base] = effort
    index_meta = {"index": raw.get("index_version") or "Artificial Analysis Intelligence Index (version not exposed by the API)",
                  "fetched": now_et().strftime("%Y-%m-%d"), "fields": used_fields, "price_label": PRICE_LABEL}
    return bench, eff, index_meta, raw


def fetch_aa_media(key, kind, probe=False):
    # include_categories asks the endpoint to say which sub-boards each model appears on. It is
    # metadata only and does not filter the rows. It is requested so the saved response records
    # whether a narrower board, such as video with audio, exists to be selected later.
    status, r = http_get(f"https://artificialanalysis.ai/api/v2/data/media/{kind}?include_categories=true",
                         {"x-api-key": key})
    if status != 200:
        raise RuntimeError(f"AA media/{kind} returned {status}")
    raw = r.json()
    models = raw.get("data", [])
    # Kept for the same reason as the llms response: when a board's numbers all move at once,
    # the only way to tell a re-grade from a different board is to read what the endpoint sent.
    os.makedirs(os.path.join(ROOT, "data", "raw"), exist_ok=True)
    with open(os.path.join(ROOT, "data", "raw", f"aa-media-{kind}.json"), "w") as f:
        json.dump(raw, f)
    if probe:
        log(f"AA media/{kind}: {len(models)} rows; keys: " + ", ".join(sorted(models[0].keys())) if models else "empty")
        return None
    d = {}
    for m in models:
        name = m.get("name") or m.get("slug")
        elo = first_key(m, ["elo", "arena_elo", "artificial_analysis_arena_elo", "score"])
        if name and elo:
            d[name] = round(float(m[elo]))
    if not d:
        raise RuntimeError(f"AA media/{kind}: no elo field found; keys were {sorted(models[0].keys()) if models else 'none'}")
    return d


# ------------------------------------------------------------------- arena.ai
def fetch_arena(board, probe=False):
    payload = fetch_public_table(ARENA_URLS[board], ROOT, "arena-" + board)
    data, excluded = parse_arena_table(payload, board)
    log(f"arena {board}: {len(data)} scores; {len(excluded)} estimated or ambiguous rows excluded")
    return data, {"url": payload["url"], "source_updated": payload["source_updated"],
                  "excluded": excluded, "board": board}


def apply_public_aa(parsed, payload, eff_rows, bench, prev, today):
    """Attach exact effort costs and fill API-missing benchmarks with row provenance."""
    notes = []
    source = {"src": "Artificial Analysis public leaderboard", "url": AA_URL,
              "fetched": today, "source_updated": payload.get("source_updated"),
              "fetched_at": payload.get("fetched_at")}
    api_names = {}
    for row in eff_rows:
        api_names.setdefault(effort_key(row.get("api_name", row["model"])), []).append(row)
    page_costs = {}
    for row in parsed.get("cost", []):
        page_costs.setdefault(effort_key(row["name"]), []).append(row)
    old_costs = {r.get("api_name", r["model"] + " (" + r["effort"] + ")"): r
                 for r in (prev or {}).get("eff", {}).get("rows", [])}
    for row in eff_rows:
        name = row.get("api_name", row["model"] + " (" + row["effort"] + ")")
        key = effort_key(name)
        matches = page_costs.get(key, [])
        match = matches[0] if len(matches) == 1 and len(api_names.get(key, [])) == 1 else None
        if match and match["value"] is not None:
            row.update(cost_per_task=match["value"], cost_display=match["display"],
                       cost_source=AA_URL, cost_fetched=today, cost_stale=False,
                       cost_status="rounded" if match["value"] == 0 else "fresh",
                       cost_join={"page_name": match["name"], "api_name": name, "key": list(key)})
        else:
            old = old_costs.get(name, {})
            if old.get("cost_per_task") is not None:
                row.update({k: v for k, v in old.items() if k.startswith("cost_")})
                row.update(cost_stale=True, cost_status="stale")
            else:
                row.update(cost_per_task=None, cost_status="missing", cost_stale=True)
            log(f"cost unavailable or ambiguous: {name}")
            if row.get("access"):
                notes.append("cost missing or carried forward")
    for key, candidates in page_costs.items():
        if key not in api_names:
            log("cost page row unjoined: " + "; ".join(r["name"] for r in candidates))
    for k in ("gdpval", "omni", "nohalluc"):
        if k in bench or k not in parsed:
            continue
        d, settings, records = {}, {}, {}
        for row in parsed[k]:
            base, setting = split_effort(row["name"])
            matches = api_names.get(effort_key(row["name"]), [])
            if len(matches) == 1:
                base, setting = matches[0]["model"], matches[0]["effort"]
            if row["value"] is None:
                log(f"{k}: missing for {row['name']}")
                continue
            if base not in d or row["value"] > d[base]:
                d[base], settings[base] = row["value"], setting
                records[base] = dict(source, page_name=row["name"], field=AA_COLUMNS[k],
                                     display=row["display"], effort=setting, stale=False)
        old = (prev or {}).get("bench", {}).get(k, {})
        stale_models = []
        for model, value in old.get("d", {}).items():
            if model not in d:
                d[model] = value
                settings[model] = old.get("settings", {}).get(model, "unknown (carried forward)")
                records[model] = dict(old.get("records", {}).get(model, {
                    "src": old.get("src"), "fetched": old.get("fetched")}), stale=True)
                stale_models.append(model)
        bench[k] = dict(source, name=BENCH[k]["name"], kind="task", elo=False,
                        field=AA_COLUMNS[k], d=d, settings=settings, records=records,
                        stale=bool(stale_models), stale_models=stale_models)
        if stale_models:
            notes.append(k + " has carried-forward model scores")
    return sorted(set(notes))


# --------------------------------------------------------------------- inbox
def parse_kv_items(path):
    """Parse the agents' '### Item N' / '### Story N' blocks of key: value lines."""
    items, cur = [], None
    with open(path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    for ln in lines:
        if ln.startswith("## "):
            cur = None
        elif ln.startswith("### "):
            cur = {"_title": ln[4:].strip()}
            items.append(cur)
        elif cur is not None and re.match(r"^\s*\d+\.\s+\S", ln):
            cur.setdefault("_links", []).append(ln.strip().split(None, 1)[1])
        elif cur is not None and ":" in ln and not ln.startswith(" "):
            k, v = ln.split(":", 1)
            cur[k.strip()] = v.strip()
    head = {}
    for ln in lines[:12]:
        if ":" in ln and not ln.startswith("#"):
            k, v = ln.split(":", 1)
            head[k.strip()] = v.strip()
    return head, items


def read_news(today):
    path = os.path.join(ROOT, "inbox", "news", f"{today}-edited.md")
    if not os.path.exists(path):
        return [], {"editor_file": None, "note": "No Editor picks on main for today. Showing yesterday's stories."}
    head, items = parse_kv_items(path)
    stories = []
    for it in items:
        if not it["_title"].lower().startswith("story"):
            continue
        links = []
        for i, u in enumerate(it.get("_links", [])):
            label = re.sub(r"^www\.", "", re.sub(r"^https?://", "", u)).split("/")[0]
            srcing = it.get("sourcing", "")
            role = "primary" if i == 0 and srcing != "vendor_only" else ("vendor" if srcing == "vendor_only" and i == 0 else "reported")
            links.append([label, u, role])
        stories.append({"tag": f"{it.get('beat', '').title()} · {it.get('status', '')}", "title": it.get("headline", it["_title"]),
                        "body": it.get("summary", ""), "why": it.get("why_it_matters", ""), "sourcing": it.get("sourcing", ""),
                        "links": links, "notes": it.get("notes", "")})
    note = f"Chosen by the Editor agent from {head.get('items_considered', '?')} items filed by the News Scout."
    try:
        if int(head.get("items_considered", "0")) < 10:
            note += " That is under the 10-item bar, so this is a thin choice set rather than a real selection."
    except ValueError:
        pass
    return stories, {"editor_file": os.path.relpath(path, ROOT), "considered": head.get("items_considered"), "note": note}


def read_scout(today):
    """Verified Scout items with a numeric score become the 'second opinions' cards."""
    bpath = os.path.join(ROOT, "inbox", "benchmarks", f"{today}.md")
    vpath = os.path.join(ROOT, "inbox", "verified", f"{today}.md")
    if not os.path.exists(bpath):
        return [], {"scout_file": None, "verified_file": None, "note": "No Benchmark Scout report on main for today."}
    _, items = parse_kv_items(bpath)
    verdicts = {}
    if os.path.exists(vpath):
        _, vitems = parse_kv_items(vpath)
        for v in vitems:
            verdicts[v["_title"]] = v.get("verdict", "unverified")
    out = []
    for it in items:
        score = it.get("score", "")
        num = re.search(r"-?\d+(?:\.\d+)?", score)
        verdict = verdicts.get(it["_title"], "unverified") if verdicts else "unverified"
        out.append({"item": it["_title"], "type": it.get("type", ""), "model": it.get("model", ""), "test": it.get("test", ""),
                    "score": score, "value": float(num.group()) if num else None, "settings": it.get("settings", ""),
                    "source": it.get("source", ""), "url": it.get("url", ""), "published": it.get("published", ""),
                    "quote": it.get("quote", ""), "verdict": verdict})
    n_pass = sum(1 for o in out if o["verdict"] == "pass")
    note = f"{len(out)} items filed by the Benchmark Scout; {n_pass} confirmed against the source page by the Verifier."
    if not verdicts:
        note = f"{len(out)} items filed by the Benchmark Scout. No Verifier report on main for today, so every item is unverified."
    return out, {"scout_file": os.path.relpath(bpath, ROOT), "verified_file": os.path.relpath(vpath, ROOT) if verdicts else None, "note": note}


def allowlist_sections():
    """Returns (allow_text, block_text) from sources/allowlist.md, lowercased."""
    p = os.path.join(ROOT, "sources", "allowlist.md")
    if not os.path.exists(p):
        return "", ""
    with open(p, encoding="utf-8") as f:
        txt = f.read()
    m = re.search(r"^## Blocklist\s*$", txt, re.M)
    if not m:
        return txt.lower(), ""
    return txt[:m.start()].lower(), txt[m.start():].lower()


def read_proposed_sources():
    """inbox/proposed-sources.md is the Scout's append-only queue. Each unresolved block
    (domain on neither list) is shown with whatever facts the Verifier has recorded for it,
    and nothing else: no rating, no recommendation. The decision is a human edit to
    sources/allowlist.md."""
    p = os.path.join(ROOT, "inbox", "proposed-sources.md")
    if not os.path.exists(p):
        return []
    allow, block = allowlist_sections()
    blocks, cur = [], None
    with open(p, encoding="utf-8") as f:
        for ln in f.read().splitlines():
            if ln.startswith("- url:"):
                cur = {"url": ln[6:].strip()}
                blocks.append(cur)
            elif cur is not None and ln.startswith("  ") and ":" in ln:
                k, v = ln.strip().split(":", 1)
                cur[k.strip()] = v.strip()
    # latest Verifier facts per domain, from the "Proposed source checks" blocks
    facts = {}
    for vp in sorted(glob.glob(os.path.join(ROOT, "inbox", "verified", "????-??-??.md"))):
        _, vitems = parse_kv_items(vp)
        for v in vitems:
            if "resolves" in v or "own_numbers" in v:
                facts[v["_title"].strip().lower()] = dict(v, checked_on=os.path.basename(vp)[:10])
    out = []
    for b in blocks:
        d = b.get("domain", "").strip().lower()
        if not d:
            continue
        status = "unresolved"
        if d in block:
            status = "blocklist"
        elif d in allow:
            status = "allowlist"
        if status != "unresolved":
            continue
        f = facts.get(d)
        out.append({"domain": d, "url": b.get("url", ""), "first_seen": b.get("first_seen", ""), "found_via": b.get("found_via", ""),
                    "why": b.get("why", ""), "checked_on": f["checked_on"] if f else None,
                    "facts": {k: f[k] for k in ("resolves", "own_numbers", "methodology", "who_runs_it", "last_updated", "also_on") if f and k in f} if f else None})
    return out


def read_commentary(today, numbers, names=()):
    """inbox/commentary/YYYY-MM-DD.md, written by the Commentator agent. Sections:
    '## picks' with '### <task id>' blocks, '## changes' with '- ' lines, '## headline'.
    Every number in the text must appear in the day's data or the block is flagged."""
    path = os.path.join(ROOT, "inbox", "commentary", f"{today}.md")
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        txt = f.read()
    out = {"file": os.path.relpath(path, ROOT), "picks": {}, "changes": [], "headline": "", "flags": []}
    section, task = None, None
    for ln in txt.splitlines():
        if ln.startswith("## "):
            section, task = ln[3:].strip().lower(), None
        elif ln.startswith("### ") and section == "picks":
            task = ln[4:].strip().lower()
        elif section == "picks" and task and ln.strip():
            out["picks"][task] = (out["picks"].get(task, "") + " " + ln.strip()).strip()
        elif section == "changes" and ln.strip().startswith("- "):
            out["changes"].append(ln.strip()[2:])
        elif section == "headline" and ln.strip():
            out["headline"] = (out["headline"] + " " + ln.strip()).strip()
    allowed = set(numbers)
    out["flagged_tasks"] = set()
    for where, text in [(f"picks/{k}", v) for k, v in out["picks"].items()] + [("changes", " ".join(out["changes"]))]:
        # model and test names carry version numbers (Fable 5.1, Terminal-Bench v2.1); those are not scores
        for nm in sorted(names, key=len, reverse=True):
            text = text.replace(nm, " ")
        text = re.sub(r"\d{4}-\d{2}-\d{2}", " ", text)
        for n in re.findall(r"\d+(?:\.\d+)?", text):
            if n not in allowed and n.rstrip("0").rstrip(".") not in allowed:
                out["flags"].append(f"{where}: {n} is not in today's data")
                if where.startswith("picks/"):
                    out["flagged_tasks"].add(where[6:])
    out["flagged_tasks"] = sorted(out["flagged_tasks"])
    return out


def known_names(bench):
    names = set()
    for b in bench.values():
        names.add(b["name"])
        names.update(b["d"].keys())
    return names


def number_set(bench, eff):
    s = set()
    for b in bench.values():
        for v in b["d"].values():
            s.add(str(v)); s.add(str(round(v))); s.add(str(v).rstrip("0").rstrip("."))
    for r in eff:
        for k in ("score", "price_1m", "speed", "wait_s"):
            v = r.get(k)
            if v is not None:
                s.add(str(v)); s.add(str(round(v))); s.add(f"{v:.2f}"); s.add(f"{v:.0f}")
    for n in range(0, 101):
        s.add(str(n))  # small integers are counts and percentages, not scores
    return s


# --------------------------------------------------------------------- picks
def normalise(b):
    if not b.get("elo"):
        return dict(b["d"])
    vals = list(b["d"].values())
    lo, hi = min(vals), max(vals)
    return {m: (v - lo) / (hi - lo) * 100 if hi > lo else 100 for m, v in b["d"].items()}


def score_task(t, bench):
    keys = list(t["w"])
    norm = {k: normalise(bench[k]) for k in keys if k in bench}
    if len(norm) < len(keys):
        return None
    models = set()
    for k in keys:
        models |= set(bench[k]["d"])
    access = MEDIA_ACCESS if t.get("media") else ACCESS
    rows = []
    for m in models:
        if all(m in norm[k] for k in keys):
            rows.append({"m": m, "s": sum(norm[k][m] * t["w"][k] for k in keys), "access": access.get(m)})
    rows.sort(key=lambda r: -r["s"])
    acc = [r for r in rows if r["access"]]
    missing = [m for m in models if access.get(m) and not any(r["m"] == m for r in rows)]
    return {"rows": rows, "pick": acc[0] if acc else None, "runner": acc[1] if len(acc) > 1 else None,
            "overall": rows[0] if rows else None, "missing": missing}


def templated_why(t, r, bench):
    """Plain sentences built from the numbers only."""
    if not r or not r["pick"]:
        return "No model you can open has a score on every test in this blend."
    p = r["pick"]
    parts = []
    if t.get("media"):
        k = list(t["w"])[0]
        board = bench[k]["d"]
        order = sorted(board.items(), key=lambda kv: -kv[1])
        rank = [m for m, _ in order].index(p["m"]) + 1
        parts.append(f"{p['m']} is ranked {rank} of {len(order)} on this board at {board[p['m']]}.")
        if rank > 1:
            parts.append(f"The board leader is {order[0][0]} at {order[0][1]}, which you cannot open.")
        return " ".join(parts)
    leads = [bench[k]["name"] for k in t["w"] if max(bench[k]["d"].items(), key=lambda kv: kv[1])[0] == p["m"]]
    if leads:
        parts.append(f"{p['m']} leads {' and '.join(leads)} outright.")
    else:
        parts.append(f"{p['m']} has the best blend across the tests without leading any one of them.")
    if r["runner"]:
        gap = p["s"] - r["runner"]["s"]
        parts.append(f"{r['runner']['m']} is {gap:.0f} points back" + (", within noise." if gap < NOISE else "."))
    if r["overall"] and r["overall"]["m"] != p["m"]:
        parts.append(f"The overall leader is {r['overall']['m']}, which you cannot open.")
    return " ".join(parts)


def compute_picks(bench, commentary):
    picks = {}
    for t in TASKS:
        r = score_task(t, bench)
        if r is None:
            picks[t["id"]] = {"name": t["name"], "pick": None, "runner": None, "overall": None, "rows": [], "missing": [],
                              "confidence": "no data", "evidence": "none", "why": "One of the tests in this blend has no data today.",
                              "why_source": "templated"}
            continue
        gap = (r["pick"]["s"] - r["runner"]["s"]) if r["pick"] and r["runner"] else 99
        kinds = [bench[k]["kind"] for k in t["w"]]
        evidence = "Preference votes" if all(k == "votes" for k in kinds) else ("Tests plus votes" if "votes" in kinds else "Task tests")
        why, src = templated_why(t, r, bench), "templated"
        if commentary and t["id"] in commentary["picks"]:
            if t["id"] in commentary["flagged_tasks"]:
                src = "flagged"  # keep the templated sentence; the page says the agent's line was set aside
            else:
                why, src = commentary["picks"][t["id"]], "commentator"
        picks[t["id"]] = {"name": t["name"], "pick": r["pick"]["m"] if r["pick"] else None,
                          "runner": r["runner"]["m"] if r["runner"] else None, "overall": r["overall"]["m"] if r["overall"] else None,
                          "rows": [[x["m"], round(x["s"], 1), x["access"]] for x in r["rows"]], "missing": r["missing"],
                          "confidence": "Within noise" if gap < NOISE else "Clear lead", "evidence": evidence,
                          "why": why, "why_source": src, "w": t["w"], "extra": t.get("extra"), "media": bool(t.get("media")),
                          "tests": t["tests"], "daily": t["daily"]}
    return picks


# ------------------------------------------------------------------- changes
# Group broad shifts without asserting an unverified explanation for their cause.
SHIFT_MIN_MODELS = 5      # below this a board is too small for "nearly every" to mean anything
SHIFT_MOVED_FRACTION = 0.7   # share of comparable models that must have moved at all
SHIFT_AGREE_FRACTION = 0.9   # share of those moves that must point the same way


def uniform_shift(b, pb):
    """One line if this whole board moved together, else None.

    Returns (text, moved, comparable, median delta) so the caller can also suppress the
    per-model lines for the same test.
    """
    deltas = {m: v - pb["d"][m] for m, v in b["d"].items() if m in pb["d"] and isinstance(pb["d"][m], (int, float))}
    if len(deltas) < SHIFT_MIN_MODELS:
        return None
    moved = [d for d in deltas.values() if abs(d) >= NOISE]
    if len(moved) < SHIFT_MOVED_FRACTION * len(deltas):
        return None
    up = sum(1 for d in moved if d > 0)
    agree = max(up, len(moved) - up)
    if agree < SHIFT_AGREE_FRACTION * len(moved):
        return None
    med = sorted(moved)[len(moved) // 2]
    direction = "up" if up >= len(moved) - up else "down"
    return (f"{b['name']} shifted across the board: {len(moved)} of {len(deltas)} models moved {direction} "
            f"by about {abs(med):.1f} at the middle. The cause is unconfirmed; this may reflect grading, "
            f"board composition, or model changes. Individual moves on this test are grouped today.",
            len(moved), len(deltas), med)



def oldest_stale_date(b):
    dates = [r.get("fetched", "unknown") for r in b.get("records", {}).values() if r.get("stale")]
    return min(dates) if dates else b.get("fetched", "unknown")


def compute_changes(today_d, prev):
    changes = []
    if not prev:
        changes.append(["mid", "First automated rebuild. There is no previous data file to compare against."])
        return changes
    flips = 0
    for tid, p in today_d["picks"].items():
        q = prev.get("picks", {}).get(tid)
        if q and q.get("pick") and p["pick"] and q["pick"] != p["pick"]:
            flips += 1
            changes.append(["pick", f"{p['name']}: pick changed from {q['pick']} to {p['pick']}. Runner-up is now {p['runner']}."])
    if flips == 0:
        changes.append(["good", "No picks changed today. Every task card carries the same model as yesterday."])
    for k, b in today_d["bench"].items():
        pb = prev.get("bench", {}).get(k)
        if not pb:
            changes.append(["good", f"New test on the page: {b['name']}."])
            continue
        shift = uniform_shift(b, pb)
        if shift:
            changes.append(["regrade", shift[0]])
        for m, v in b["d"].items():
            if m not in pb["d"]:
                if ACCESS.get(m) or MEDIA_ACCESS.get(m):
                    changes.append(["good", f"New on {b['name']}: {m} at {v}."])
            elif shift:
                continue  # the board moved as a whole; a per-model line would misread it as progress
            elif (ACCESS.get(m) or MEDIA_ACCESS.get(m)) and abs(v - pb["d"][m]) >= NOISE:
                changes.append(["mid", f"{b['name']}: {m} moved from {pb['d'][m]} to {v}."])
    stale = [b for b in today_d["bench"].values() if b.get("stale")]
    if len(stale) > 3:
        changes.append(["weak", f"{len(stale)} of {len(today_d['bench'])} sources include carried-forward scores: "
                        + ", ".join(f"{b['name']} (oldest carried-forward score: {oldest_stale_date(b)})" for b in stale) + "."])
    else:
        for b in stale:
            changes.append(["weak", f"{b['name']} includes old scores from {oldest_stale_date(b)}; see source records for each model."])
    return changes


# -------------------------------------------------------------------- health
def read_runs():
    rows = []
    for a in AGENTS:
        p = os.path.join(ROOT, "log", f"runs-{a}.csv")
        if not os.path.exists(p):
            continue
        with open(p, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r.get("date"):
                    rows.append(r)
    rows.sort(key=lambda r: (r["date"], r.get("time_et", "")))
    return rows


def health_strip(rows, today, days=14):
    cutoff = (dt.date.fromisoformat(today) - dt.timedelta(days=days)).isoformat()
    out = [r for r in rows if r["date"] >= cutoff]
    seen_today = {r["agent"] for r in out if r["date"] == today}
    for a in AGENTS:
        if a != "rebuild" and a not in seen_today:
            out.append({"date": today, "time_et": "", "agent": a, "status": "did not run", "items": "0", "pr": "", "notes": "no log line for today"})
    return [{"agent": AGENT_LABEL.get(r["agent"], r["agent"]), "date": r["date"], "time": r.get("time_et", ""), "status": r["status"],
             "items": r.get("items", ""), "pr": r.get("pr", ""), "notes": r.get("notes", "")} for r in out]


def write_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    header = ["date", "time_et", "agent", "status", "items", "pr", "notes"]
    combined = wb.active
    combined.title = "combined"
    combined.append(header)
    for r in rows:
        combined.append([r.get(h, "") for h in header])
    for a in AGENTS:
        ws = wb.create_sheet(a[:31])
        ws.append(header)
        for r in rows:
            if r["agent"] == a:
                ws.append([r.get(h, "") for h in header])
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col, width in zip("ABCDEFG", (12, 9, 18, 10, 7, 6, 60)):
            ws.column_dimensions[col].width = width
    os.makedirs(os.path.join(ROOT, "log"), exist_ok=True)
    wb.save(os.path.join(ROOT, "log", "run-log.xlsx"))


def append_log(today, status, items, pr, notes):
    p = os.path.join(ROOT, "log", "runs-rebuild.csv")
    notes = notes.replace(",", ";")[:99]
    line = f"{today},{now_et().strftime('%H:%M')},rebuild,{status},{items},{pr},{notes}\n"
    if not os.path.exists(p):
        with open(p, "w") as f:
            f.write("date,time_et,agent,status,items,pr,notes\n")
    with open(p, "a") as f:
        f.write(line)
    log("log line: " + line.strip())


# ------------------------------------------------------------- pending PRs
def pending_prs():
    try:
        import requests
        h = {"Accept": "application/vnd.github+json"}
        tok = os.environ.get("GITHUB_TOKEN")
        if tok:
            h["Authorization"] = f"Bearer {tok}"
        repo = os.environ.get("GITHUB_REPOSITORY", "jwatson7399/ai-auditing-dashboard")
        r = requests.get(f"https://api.github.com/repos/{repo}/pulls?state=open&per_page=30", headers=h, timeout=30)
        if r.status_code != 200:
            return [{"number": 0, "title": f"could not list pull requests ({r.status_code})"}]
        out = []
        for p in r.json():
            fr = requests.get(p["url"] + "/files", headers=h, timeout=30)
            files = [f["filename"] for f in fr.json()] if fr.status_code == 200 else []
            if any(f.startswith("inbox/") for f in files):
                out.append({"number": p["number"], "title": p["title"], "files": [f for f in files if f.startswith("inbox/")]})
        return out
    except Exception as e:  # noqa: BLE001
        return [{"number": 0, "title": f"could not list pull requests ({e.__class__.__name__})"}]


# ---------------------------------------------------------------------- page
def build_page(d):
    tpl = os.path.join(ROOT, "site", "template.html")
    with open(tpl, encoding="utf-8") as f:
        html = f.read()
    if "\u2014" in html:
        raise RuntimeError("template.html contains an em dash; the template is house text and must not")
    payload = json.dumps(d, ensure_ascii=False).replace("</", "<\\/")
    html = html.replace("/*DATA*/null", payload)
    os.makedirs(os.path.join(ROOT, "site"), exist_ok=True)
    with open(os.path.join(ROOT, "site", "index.html"), "w", encoding="utf-8") as f:
        f.write(html)


# ---------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", default=now_et().strftime("%Y-%m-%d"))
    ap.add_argument("--no-fetch", action="store_true", help="reuse the last data file's numbers")
    ap.add_argument("--probe", action="store_true", help="print what each source returns and exit")
    ap.add_argument("--no-log", action="store_true")
    ap.add_argument("--skip-arena", action="store_true")
    args = ap.parse_args()
    today = args.date
    key = os.environ.get("AA_API_KEY", "").strip()
    run_id = os.environ.get("GITHUB_RUN_NUMBER", "local")

    if args.probe:
        try:
            fetch_aa_llms(key, probe=True)
            for kind in ("text-to-image", "text-to-video"):
                try:
                    fetch_aa_media(key, kind, probe=True)
                except Exception as e:  # noqa: BLE001
                    log(f"media/{kind}: {e}")
            if not args.skip_arena:
                for board in ("text", "webdev"):
                    try:
                        fetch_arena(board, probe=True)
                    except Exception as e:  # noqa: BLE001
                        log(f"arena {board}: {e}")
        except Blocked as e:
            log(f"blocked: {e}")
        return 0

    status, notes, items = "ok", [], 0
    prev = load_previous(today)
    carry = prev
    latest = os.path.join(ROOT, "data", "latest.json")
    if os.path.exists(latest):
        with open(latest) as f:
            candidate = json.load(f)
        if candidate.get("date", "9999") <= today and (not carry or candidate["date"] >= carry["date"]):
            carry = candidate  # retries retain today's latest good scores; diffs still use yesterday
    try:
        # 1. numbers
        bench = {}
        eff_rows, eff_meta = [], {}
        fresh = 0
        if args.no_fetch:
            same = os.path.join(ROOT, "data", f"{today}.json")
            if not prev and os.path.exists(same):
                with open(same) as f:
                    prev = json.load(f)  # rebuild today's page from today's own numbers
            if not prev:
                raise Blocked("--no-fetch with no previous data file")
            bench = {k: dict(v, stale=True) for k, v in prev["bench"].items()}
            for b in bench.values():
                b["records"] = {m: dict(r, stale=True) for m, r in b.get("records", {}).items()}
            eff_rows, eff_meta = prev["eff"]["rows"], {k: v for k, v in prev["eff"].items() if k != "rows"}
            for row in eff_rows:
                if row.get("cost_per_task") is not None:
                    row.update(cost_stale=True, cost_status="stale")
            notes.append("numbers carried forward (--no-fetch)")
            status = "partial"
        else:
            got, eff_rows, eff_meta, raw = fetch_aa_llms(key)
            os.makedirs(os.path.join(ROOT, "data", "raw"), exist_ok=True)
            with open(os.path.join(ROOT, "data", "raw", "aa-llms-latest.json"), "w") as f:
                json.dump(raw, f)
            for k, b in BENCH.items():
                if k in got and got[k]["d"]:
                    bench[k] = {"name": b["name"], "kind": b["kind"], "elo": bool(b.get("elo")), "src": "Artificial Analysis",
                                "fetched": today, "field": eff_meta["fields"].get(k), "index_version": eff_meta["index"],
                                "stale": False, "d": got[k]["d"], "settings": got[k]["settings"]}
                    fresh += 1
            try:
                public = fetch_public_table(AA_URL, ROOT, "aa-leaderboard", expand=True)
                parsed, errors = parse_aa_table(public["rows"])
                for field, error in errors.items():
                    log(f"AA public {field}: {error}")
                if errors:
                    notes.append("AA public columns unavailable: " + ", ".join(errors))
                    status = "partial"
            except Exception as e:
                log(f"AA public leaderboard: {e}")
                parsed, public = {}, {}
                notes.append("AA public leaderboard unavailable")
                status = "partial"
            public_notes = apply_public_aa(parsed, public, eff_rows, bench, carry, today)
            notes.extend(public_notes)
            if public_notes:
                status = "partial"
            fresh = sum(not b.get("stale") for b in bench.values())
            for k, b in BENCH.items():
                if k in bench:
                    continue
                try:
                    source_meta = {}
                    if b.get("media"):
                        d = fetch_aa_media(key, b["media"])
                        src = "Artificial Analysis"
                    elif b.get("arena") and not args.skip_arena:
                        d, source_meta = fetch_arena(b["arena"])
                        src = "arena.ai"
                    else:
                        raise RuntimeError("no fetcher produced this benchmark")
                    # These two endpoints name models differently from the llms endpoint, so the
                    # names are folded onto the access table before anything scores them.
                    old = (carry or {}).get("bench", {}).get(k, {})
                    names = dict.fromkeys(old.get("d", {})) if b.get("arena") else {}
                    names.update(MEDIA_ACCESS if b.get("media") else ACCESS)
                    d, aliases, unmatched = fold_to_access(d, names)
                    stale_models = []
                    records = {m: {"fetched": today, "src": src, "stale": False,
                                   "url": source_meta.get("url"), "source_updated": source_meta.get("source_updated"),
                                   "raw_names": [raw for raw, base in aliases.items() if base == m] or [m]}
                               for m in d}
                    old = (carry or {}).get("bench", {}).get(k, {})
                    for model, value in old.get("d", {}).items():
                        if model not in d:
                            d[model] = value
                            records[model] = dict(old.get("records", {}).get(model, {
                                "fetched": old.get("fetched"), "src": old.get("src")}), stale=True)
                            stale_models.append(model)
                    if stale_models:
                        notes.append(k + " has carried-forward model scores")
                        status = "partial"
                    if aliases:
                        log(f"{k}: folded " + "; ".join(f"{r} -> {v}" for r, v in sorted(aliases.items())))
                    bench[k] = {"name": b["name"], "kind": b["kind"], "elo": bool(b.get("elo")), "src": src, "fetched": today,
                                "stale": bool(stale_models), "stale_models": stale_models, "records": records,
                                "d": d, "aliases": aliases, "unmatched": unmatched, **source_meta}
                    fresh += 1
                except Exception as e:  # noqa: BLE001
                    log(f"{k}: {e}")
                    if carry and k in carry["bench"]:
                        bench[k] = dict(carry["bench"][k], stale=True)
                        bench[k]["records"] = {m: dict(r, stale=True) for m, r in bench[k].get("records", {}).items()}
                        notes.append(f"{k} carried forward")
                    else:
                        notes.append(f"{k} missing")
                    status = "partial"
            items = fresh
        if not bench:
            raise Blocked("no benchmark data at all")

        # 2. inbox on main
        news, news_meta = read_news(today)
        if not news and prev:
            news, news_meta = prev.get("news", []), dict(prev.get("news_meta", {}), note=news_meta["note"])
        second, second_meta = read_scout(today)
        proposed = read_proposed_sources()
        prs = pending_prs()

        # 3. picks, commentary, changes
        commentary = read_commentary(today, number_set(bench, eff_rows), known_names(bench))
        picks = compute_picks(bench, commentary)
        d = {"schema": 1, "date": today, "built_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
             "built_by": f"tools/rebuild.py, GitHub Actions run {run_id}", "built_at_et": now_et().strftime("%-I:%M %p"),
             "status": status, "status_notes": notes,
             "access": ACCESS, "media_access": MEDIA_ACCESS, "bench": bench, "second": second, "second_meta": second_meta,
             "news": news, "news_meta": news_meta, "eff": dict(eff_meta, rows=eff_rows), "picks": picks,
             "commentary": commentary, "proposed_sources": proposed, "pending_prs": prs, "previous_date": prev["date"] if prev else None, "noise": NOISE}
        d["changes"] = compute_changes(d, prev)
        if commentary and commentary["changes"]:
            d["changes"] = [["commentator", c] for c in commentary["changes"]] + d["changes"]

        # 4. health strip and workbook
        runs = read_runs()
        d["health"] = health_strip(runs, today)
        write_xlsx(runs)

        # 5. write data and page
        os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
        with open(os.path.join(ROOT, "data", f"{today}.json"), "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=1)
        with open(os.path.join(ROOT, "data", "latest.json"), "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=1)
        build_page(d)
        log(f"built {today}: status {status}; picks " + ", ".join(f"{k}={v['pick']}" for k, v in picks.items()))
        rc = 0
    except Blocked as e:
        status, rc = "blocked", 2
        notes = [str(e)]
        log(f"BLOCKED: {e}")
    except Exception as e:  # noqa: BLE001
        status, rc = "error", 1
        notes = [f"{e.__class__.__name__}: {e}"]
        traceback.print_exc()
    finally:
        if not args.no_log:
            append_log(today, status, items, run_id, "; ".join(notes) if notes else "none")
    return rc


if __name__ == "__main__":
    sys.exit(main())
